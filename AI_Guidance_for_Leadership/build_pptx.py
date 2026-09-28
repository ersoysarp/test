"""Generate an editable PowerPoint version of the leadership deck.

Mirrors index.html: 10 slides, 16:9, Arial, same palette and structure.
Run: python3 build_pptx.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PX = 914400 / 96  # one CSS pixel in EMU

INK = RGBColor(0x15, 0x25, 0x30)
INK_SOFT = RGBColor(0x41, 0x53, 0x5D)
INK_FAINT = RGBColor(0x78, 0x88, 0x92)
PAPER = RGBColor(0xF7, 0xF5, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AZURE = RGBColor(0x0E, 0x6E, 0x8C)
AZURE_DEEP = RGBColor(0x09, 0x3F, 0x52)
GOLD = RGBColor(0xA6, 0x7C, 0x3D)
GOLD_SOFT = RGBColor(0xC7, 0x9A, 0x55)
LINE = RGBColor(0xE0, 0xDA, 0xCD)
LINE_STRONG = RGBColor(0xCF, 0xC7, 0xB5)
TINT = RGBColor(0xE9, 0xF1, 0xF4)
TINT_GOLD = RGBColor(0xF6, 0xEF, 0xE0)
SAFE = RGBColor(0x2F, 0x7A, 0x5B)
CAUTION = RGBColor(0x9A, 0x73, 0x26)
STOP = RGBColor(0x9C, 0x4A, 0x3C)
COVER_BG = RGBColor(0x0A, 0x3F, 0x52)

FONT = "Arial"


def px(v):
    return Emu(int(v * PX))


def new_slide(prs, bg=PAPER, accent=AZURE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bgshape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(1280), px(720))
    bgshape.fill.solid()
    bgshape.fill.fore_color.rgb = bg
    bgshape.line.fill.background()
    bgshape.shadow.inherit = False
    if accent is not None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(6), px(720))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()
        bar.shadow.inherit = False
    return slide


def text(slide, x, y, w, h, runs, size=12, color=INK_SOFT, bold=False,
         italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.0, caps=False):
    """runs: a string, or list of (text, {overrides}) tuples."""
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    if isinstance(runs, str):
        runs = [(runs, {})]
    for content, over in runs:
        r = p.add_run()
        r.text = content.upper() if caps else content
        f = r.font
        f.name = FONT
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.italic = over.get("italic", italic)
        f.color.rgb = over.get("color", color)
    return box


def para(box, runs, size=12, color=INK_SOFT, bold=False, italic=False,
         space_before=0, spacing=1.0, align=PP_ALIGN.LEFT):
    p = box.text_frame.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    p.space_before = Pt(space_before)
    if isinstance(runs, str):
        runs = [(runs, {})]
    for content, over in runs:
        r = p.add_run()
        r.text = content
        f = r.font
        f.name = FONT
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.italic = over.get("italic", italic)
        f.color.rgb = over.get("color", color)
    return p


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, rounded=False, line_w=1):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        px(x), px(y), px(w), px(h))
    if rounded:
        shape.adjustments[0] = 0.06
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def header(slide, num, kicker, title, accent=GOLD):
    if num:
        text(slide, 48, 34, 60, 16, num, size=10.5, color=GOLD, bold=True)
        text(slide, 88, 34, 700, 16, kicker, size=10.5, color=AZURE_DEEP,
             bold=True, caps=True)
    else:
        text(slide, 48, 34, 700, 16, kicker, size=10.5, color=GOLD,
             bold=True, caps=True)
    text(slide, 48, 56, 900, 60, title, size=18.75, color=INK, bold=True,
         spacing=1.15)
    rect(slide, 48, 100, 54, 3, fill=accent, line=None)


def footer(slide, left, right):
    rect(slide, 48, 664, 1184, 1, fill=LINE, line=None)
    text(slide, 48, 672, 900, 16, left, size=10.5, color=INK_FAINT)
    text(slide, 1000, 672, 232, 16, right, size=10.5, color=INK_FAINT,
         align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, items, size=11.25, gap=15.5, dot=GOLD):
    for i, item in enumerate(items):
        cy = y + i * gap
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(cy + 4.5),
                                   px(4.5), px(4.5))
        d.fill.solid()
        d.fill.fore_color.rgb = dot
        d.line.fill.background()
        d.shadow.inherit = False
        text(slide, x + 11, cy, w - 11, 14, item, size=size, color=INK_SOFT)


def build():
    prs = Presentation()
    prs.slide_width = px(1280)
    prs.slide_height = px(720)

    # ---------- 1 COVER ----------
    s = new_slide(prs, bg=COVER_BG, accent=None)
    text(s, 900, 48, 332, 60,
         "Prepared for the President\nof the Republic of Azerbaijan\nBaku, 2026",
         size=11.25, color=RGBColor(0x9F, 0xC2, 0xCE), align=PP_ALIGN.RIGHT,
         spacing=1.5)
    text(s, 72, 210, 900, 20, "Artificial Intelligence in Government",
         size=12, color=RGBColor(0x8F, 0xD0, 0xDF), bold=True, caps=True)
    text(s, 72, 248, 900, 100,
         [("AI in the service of ", {}), ("government", {"color": RGBColor(0xBF, 0xE6, 0xF0)})],
         size=42, color=WHITE, bold=True, spacing=1.05)
    rect(s, 72, 372, 4, 62, fill=GOLD, line=None)
    text(s, 92, 374, 700, 60,
         "Where AI creates value for the Administration today, how to protect "
         "official information, and how to build a fully secure national "
         "capability for Azerbaijan.",
         size=14, color=RGBColor(0xD3, 0xE5, 0xEC), spacing=1.45)
    rect(s, 0, 600, 1280, 120, fill=RGBColor(0x0C, 0x46, 0x5A), line=None)
    acts = [("I", "Use it safely today"),
            ("II", "Match data to the right environment"),
            ("III", "Build the sovereign capability")]
    for i, (n, label) in enumerate(acts):
        x = 72 + i * 390
        text(s, x, 640, 40, 26, n, size=19.5, color=RGBColor(0x8F, 0xD0, 0xDF), bold=True)
        text(s, x + 40, 644, 330, 40, label, size=12.75, color=WHITE, bold=True,
             spacing=1.25)

    # ---------- 2 AGENDA ----------
    s = new_slide(prs)
    header(s, None, "Agenda", "Seven questions, answered in order")
    items = [
        ("01", "Understanding AI today", "What it does well, and where judgment stays human"),
        ("02", "Three ways to use AI", "Personal, enterprise, and fully secure on premise compared"),
        ("03", "Matching data to the environment", "Azerbaijani government data against the two governed options"),
        ("04", "Deep dive: enterprise account", "Common use cases, example prompts, and proven impact"),
        ("05", "Deep dive: fully secure environment", "What it unlocks, with case studies from around the world"),
        ("06", "Building the sovereign capability", "The stack and the enablers a secure environment requires"),
        ("07", "Where to start", "A sequenced path from first use to full sovereignty"),
    ]
    y = 132
    for num, title, desc in items:
        rect(s, 48, y, 1184, 1, fill=LINE, line=None)
        text(s, 48, y + 22, 50, 20, num, size=16.5, color=GOLD, bold=True)
        text(s, 104, y + 22, 340, 24, title, size=15, color=INK, bold=True)
        text(s, 460, y + 24, 700, 20, desc, size=12, color=INK_SOFT)
        y += 70
    rect(s, 48, y, 1184, 1, fill=LINE, line=None)
    footer(s, "Sources OECD 2026, UK GDS 2025, deployment case studies",
           "Prepared for the President")

    # ---------- 3 UNDERSTANDING AI ----------
    s = new_slide(prs)
    header(s, "01", "Understanding AI today",
           "A capable assistant for drafting, analysis and visuals, not a decision maker")
    panels = [
        (48, SAFE, RGBColor(0xE6, 0xF1, 0xEB), "Where it is strong",
         "Drafting, analysis, visuals, coordination",
         [("Draft", "Speeches, congratulatory letters, press statements"),
          ("Summarize", "A 50 page ministry report into one page"),
          ("Translate", "Azerbaijani, English, Russian, Turkish"),
          ("Analyze", "Trends and figures across ministry data"),
          ("Visualize", "Charts, dashboards, one page reports"),
          ("Assist", "Agendas, directives, briefing packs")]),
        (664, STOP, RGBColor(0xF4, 0xE7, 0xE3), "Where judgment stays human",
         "It assists, it does not decide",
         [("Truth", "Confident, and sometimes simply wrong"),
          ("Accountability", "Holds no authority, carries no responsibility"),
          ("Verification", "Every output checked against a trusted source"),
          ("Sensitivity", "Weak on policy nuance and hard cases"),
          ("Disclosure", "Flag where AI shaped a public output"),
          ("Decision", "Leadership signs the decision, not the model")]),
    ]
    for x, accent, head_bg, ptitle, psub, rows in panels:
        rect(s, x, 120, 568, 424, fill=WHITE, line=LINE, rounded=True)
        rect(s, x, 120, 568, 4, fill=accent, line=None)
        rect(s, x + 1, 124, 566, 52, fill=head_bg, line=None)
        text(s, x + 22, 134, 500, 22, ptitle, size=16.5, color=INK, bold=True)
        text(s, x + 22, 156, 500, 16, psub, size=11.25, color=INK_FAINT)
        ry = 188
        for label, desc in rows:
            text(s, x + 22, ry, 150, 18, label, size=13.5, color=INK, bold=True)
            text(s, x + 170, ry + 1, 380, 32, desc, size=12, color=INK_SOFT, spacing=1.2)
            ry += 58
            if ry < 540:
                rect(s, x + 22, ry - 12, 524, 1, fill=LINE, line=None)
    rect(s, 48, 560, 1184, 46, fill=AZURE_DEEP, line=None, rounded=True)
    text(s, 70, 574, 1140, 22,
         [("Rule of thumb. ", {"color": GOLD_SOFT, "bold": True}),
          ("Let AI produce the first draft. People own the check, the decision, and the record.", {"color": WHITE})],
         size=12.75)
    footer(s, "Capture the speed of language work, protect the judgment, verification and accountability",
           "Section 01")

    # ---------- 4 THREE WAYS ----------
    s = new_slide(prs)
    header(s, "02", "Three ways to use AI",
           "The door you choose decides which data you may use")
    cols = [300, 150, 190, 190, 190, 164]
    heads = ["Environment", "Official work", "Trains on your data",
             "Control & audit", "Where data sits", "Best used for"]
    rows = [
        (STOP, "Personal account", "Public tool, private login",
         ("No", STOP), "Possibly, to train the model", "None",
         "Public cloud\nOutside national control", "Public information only\nGeneral reference, no official data"),
        (AZURE, "Enterprise account", "Contracted, administered",
         ("With controls", CAUTION), "No", "Central, full audit trail",
         "Vendor cloud\nUnder contract terms", "Public and internal\nSpeeches, strategies, open data"),
        (GOLD, "Fully secure, on premise", "Closed state platform",
         ("Yes, fully", SAFE), "No", "State owned, accredited",
         "State premises\nNever leaves Azerbaijan", "Sensitive and restricted\nCitizen appeals, cabinet briefings"),
    ]
    x = 48
    rect(s, 48, 120, 1184, 44, fill=AZURE_DEEP, line=None)
    rect(s, 48, 120, 300, 44, fill=INK, line=None)
    for w, h in zip(cols, heads):
        text(s, x + 16, 134, w - 24, 20, h, size=12, color=WHITE, bold=True)
        x += w
    y = 164
    for i, r in enumerate(rows):
        accent, name, sub, (verdict, vcol), trains, audit, sits, best = r
        rh = 136
        rect(s, 48, y, 1184, rh, fill=WHITE, line=LINE)
        rect(s, 48, y, 300, rh, fill=TINT, line=LINE)
        text(s, 68, y + 42, 260, 24, name, size=14.25, color=INK, bold=True)
        text(s, 68, y + 68, 260, 18, sub, size=10.5, color=INK_FAINT)
        cx = 348
        vals = [(verdict, vcol, True), (trains, INK_SOFT, False),
                (audit, INK_SOFT, False), (sits, INK_SOFT, False),
                (best, INK_SOFT, False)]
        for (val, col, bold), w in zip(vals, cols[1:]):
            lines = val.split("\n")
            b = text(s, cx + 16, y + 42, w - 26, 60, lines[0], size=12,
                     color=col, bold=bold, spacing=1.2)
            if len(lines) > 1:
                para(b, lines[1], size=10.5, color=INK_FAINT, spacing=1.2,
                     space_before=3)
            cx += w
        y += rh
    rect(s, 48, y + 10, 1184, 54, fill=AZURE_DEEP, line=None, rounded=True)
    text(s, 70, y + 21, 1140, 34,
         [("Where to begin. ", {"color": GOLD_SOFT, "bold": True}),
          ("The enterprise account delivers most of the value now, at a fraction of the cost of a secure build. Start there, and plan the secure platform in parallel.", {"color": WHITE})],
         size=12)
    footer(s, "Adoption begins with a governed enterprise account, not a consumer chatbot",
           "Section 02")

    # ---------- 5 MATRIX ----------
    s = new_slide(prs)
    header(s, "03", "Matching data to the environment",
           "Classification decides which environment fits each task")
    rect(s, 48, 120, 1184, 44, fill=AZURE_DEEP, line=None)
    rect(s, 48, 120, 330, 44, fill=INK, line=None)
    text(s, 68, 134, 300, 20, "Government data type", size=12, color=WHITE, bold=True)
    for cx, title, sub in [(378, "Enterprise account", "Vendor cloud, governed"),
                           (805, "Fully secure, on premise", "Closed, national control")]:
        b = text(s, cx + 18, 128, 380, 20, title, size=12, color=WHITE, bold=True)
        para(b, sub, size=10.5, color=RGBColor(0xBC, 0xD8, 0xE1))
    mrows = [
        ("Public", ["State Statistics", "Azerbaijan 2030", "COP29 speeches"],
         ("Safe", SAFE), ["Translate a COP29 address for the world",
                          "Compare Azerbaijan 2030 to OECD practice"],
         ("Safe", SAFE), ["Available, rarely needed here"]),
        ("Internal", ["Ministry reports", "Draft policy", "Process notes"],
         ("With controls", CAUTION), ["Summarize ministry reports, draft policy",
                                      "Design cross agency trackers"],
         ("Safe", SAFE), ["The natural home, no external processing"]),
        ("Limited access", ["Citizen appeals", "Personal records", "Draft orders"],
         ("Not permitted", STOP), ["Personal data must stay in national control"],
         ("Safe", SAFE), ["Summarize citizen appeals, route to ministries",
                          "Draft replies citing the relevant law"]),
        ("Restricted", ["Security assessments", "Security Council minutes"],
         ("Not permitted", STOP), ["Never outside an accredited environment"],
         ("Safe", SAFE), ["Cross domain analysis for the Security Council",
                          "Leadership dashboards and briefing packs"]),
    ]
    y = 164
    for tier, chips, (v1, c1), u1, (v2, c2), u2 in mrows:
        rh = 104
        rect(s, 48, y, 1184, rh, fill=WHITE, line=LINE)
        rect(s, 48, y, 330, rh, fill=TINT, line=LINE)
        text(s, 68, y + 14, 290, 22, tier, size=14.25, color=INK, bold=True)
        chx, chy = 68, y + 44
        for chip in chips:
            w = 18 + 6.9 * len(chip)
            if chx + w > 366:
                chx, chy = 68, chy + 27
            rect(s, chx, chy, w, 23, fill=WHITE, line=LINE_STRONG, rounded=True)
            cb = text(s, chx + 9, chy + 5, w - 18, 16, chip, size=10.5,
                      color=AZURE_DEEP, bold=True)
            cb.text_frame.word_wrap = False
            chx += w + 6
        for cx, verdict, col, uses in [(378, v1, c1, u1), (805, v2, c2, u2)]:
            d = s.shapes.add_shape(MSO_SHAPE.OVAL, px(cx + 18), px(y + 19),
                                   px(8), px(8))
            d.fill.solid(); d.fill.fore_color.rgb = col
            d.line.fill.background(); d.shadow.inherit = False
            text(s, cx + 32, y + 14, 300, 18, verdict, size=12, color=col, bold=True)
            bullets(s, cx + 18, y + 44, 400, uses, size=11.25, gap=22)
        y += rh
    rect(s, 48, y + 8, 1184, 54, fill=AZURE_DEEP, line=None, rounded=True)
    text(s, 70, y + 20, 1140, 34,
         [("Read it this way. ", {"color": GOLD_SOFT, "bold": True}),
          ("Most everyday work is Public or Internal, so adoption can begin now. The sensitive tiers, where the largest gains sit, wait for the secure platform.", {"color": WHITE})],
         size=12)
    footer(s, "Where a task mixes tiers, the highest tier present governs the environment",
           "Section 03")

    # ---------- 6 & 7 DEEP DIVES ----------
    def deep_dive(kicker_num, kicker, title, boxes, gold=False, proof=None,
                  foot_left="", foot_right=""):
        s = new_slide(prs, accent=GOLD if gold else AZURE)
        header(s, kicker_num, kicker, title, accent=GOLD if gold else GOLD)
        accent = GOLD if gold else AZURE
        tint = TINT_GOLD if gold else TINT
        bh = 220 if proof else 240
        for i, (btitle, items, prompt) in enumerate(boxes):
            bx = 48 + (i % 2) * 604
            by = 122 + (i // 2) * (bh + 14)
            rect(s, bx, by, 580, bh, fill=WHITE, line=LINE, rounded=True)
            rect(s, bx, by, 580, 3, fill=accent, line=None)
            text(s, bx + 18, by + 14, 540, 22, btitle, size=13.5, color=INK, bold=True)
            bullets(s, bx + 18, by + 44, 540, items, size=11.25, gap=18,
                    dot=GOLD)
            ph = 70
            py0 = by + bh - ph - 10
            rect(s, bx + 14, py0, 552, ph, fill=tint, line=None, rounded=True)
            text(s, bx + 26, py0 + 9, 530, 14, "Example prompt",
                 size=10.5, color=GOLD if gold else AZURE_DEEP, bold=True, caps=True)
            text(s, bx + 26, py0 + 27, 528, 40, prompt, size=10.5,
                 color=INK_SOFT, italic=True, spacing=1.25)
        if proof:
            py = 122 + 2 * (bh + 14)
            rect(s, 48, py, 1184, 60, fill=TINT, line=LINE, rounded=True)
            b = text(s, 68, py + 13, 380, 18, proof[0], size=12,
                     color=AZURE_DEEP, bold=True)
            para(b, proof[1], size=10.5, color=INK_SOFT)
            cx = 452
            for n, lab in proof[2]:
                text(s, cx, py + 18, 92, 26, n, size=18, color=AZURE_DEEP, bold=True)
                text(s, cx + 82, py + 24, 110, 18, lab, size=11.25, color=INK_SOFT)
                cx += 190
        footer(s, foot_left, foot_right)

    deep_dive("04", "Deep dive, enterprise account",
              "Public and internal work: use cases, each with a prompt",
              [("Draft & translate",
                ["Speeches, remarks and toasts",
                 "Congratulatory and condolence letters",
                 "Press statements and talking points",
                 "Translate to English, Russian, Turkish",
                 "Polish tone and register"],
                "Translate this cleared statement into formal English. Keep official titles exact."),
               ("Summarize & analyze",
                ["Condense 50 page reports to one page",
                 "Compare policy to OECD and EU practice",
                 "Spot trends across ministry data",
                 "Benchmark against peer states",
                 "Extract action points from long files"],
                "Compare Azerbaijan 2030 to OECD practice. Return a table: our approach, theirs, the gap."),
               ("Visualize & report",
                ["Charts and graphs from spreadsheets",
                 "One page visual reports",
                 "Slide packs from State Statistics data",
                 "Simple decision dashboards",
                 "Infographics for the public"],
                "Turn this budget spreadsheet into a one page visual: a chart per program, three takeaways."),
               ("Assist & coordinate",
                ["Draft agendas and minutes",
                 "Assemble briefing packs",
                 "Track directives and actions",
                 "Prepare meeting questions",
                 "Reminders and follow ups"],
                "From these papers, draft a chair note: the decision needed and one question per item.")],
              proof=("Proven, United Kingdom",
                     "M365 Copilot, 20,000 officials, 12 departments, 3 months",
                     [("26 min", "saved daily"), ("85%", "good value"),
                      ("82%", "would keep it"), ("80%", "sustained use")]),
              foot_left="Departments included tax, welfare and justice; governance and data access came first. Source UK GDS 2025",
              foot_right="Section 04")

    deep_dive("05", "Deep dive, fully secure environment",
              "On sensitive data: use cases, each with a prompt",
              [("Citizen appeals & correspondence",
                ["Triage and summarize appeals",
                 "Route to the responsible ministry",
                 "Draft replies citing the relevant law",
                 "Detect duplicates and patterns",
                 "Track status through to closure"],
                "Summarize this appeal, name the responsible ministry, draft a reply citing the articles of law."),
               ("Leadership briefings & dashboards",
                ["Cross ministry briefing packs",
                 "Live decision dashboards",
                 "Daily situation updates",
                 "Track KPIs across domains",
                 "Visual summaries for the President"],
                "Build a dashboard from these inputs: one chart per priority, status and top risk for each."),
               ("Security & defense",
                ["Cross domain analysis for the Council",
                 "Risk maps by region",
                 "Scenario and readiness planning",
                 "Synthesize classified reports",
                 "Recommended actions with evidence"],
                "Across these reports, produce a risk map by region with evidence and one action each."),
               ("Reconstruction, Great Return",
                ["Delivery dashboards by district",
                 "Budget and progress tracking",
                 "Blocker and delay detection",
                 "Contractor and milestone monitoring",
                 "Status reports for leadership"],
                "Build a status dashboard by district: budget, progress, blockers, next milestone.")],
              gold=True,
              foot_left="All processing stays inside the accredited platform, in national control",
              foot_right="Section 05")

    # ---------- 8 CASE STUDIES ----------
    s = new_slide(prs, accent=GOLD)
    header(s, "05", "Deep dive, fully secure environment",
           "The measured impact of these cases, proven worldwide")
    cases = [
        ("Defense", "National force planning", "10x",
         "Digital twin makes force gaps transparent. 2x faster mission planning, 5x budget rationalization"),
        ("Public services", "ASAN style delivery", "1,300h",
         "Regulatory synthesizer saves hours each month. Assessments cut from 40 to 20 minutes"),
        ("Major ministry", "30,000+ staff", "1.4mn h",
         "Initiative tracker saves hours a year, and lifts citizen satisfaction by 9%"),
        ("Citizen appeals", "Replies to the President", "75%",
         "Drafting time cut. 95% of drafts need no edit, 100% of relevant laws identified"),
        ("President's Office", "Executive assistant", "40%",
         "Fewer meetings. Briefing packs and directive tracking, up to 2 hours a day saved"),
        ("Leadership analytics", "Dashboards for the President", "Minutes",
         "One set of numbers across ministries, answered on demand at the point of decision"),
    ]
    y = 122
    for name, sub, metric, out in cases:
        rh = 66
        rect(s, 48, y, 1184, rh, fill=WHITE, line=LINE, rounded=True)
        rect(s, 48, y, 250, rh, fill=TINT, line=None)
        text(s, 68, y + 16, 220, 20, name, size=13.5, color=INK, bold=True)
        text(s, 68, y + 38, 220, 16, sub, size=10.5, color=INK_FAINT)
        text(s, 318, y + 18, 150, 32, metric, size=21, color=AZURE_DEEP, bold=True)
        text(s, 470, y + 22, 740, 32, out, size=12, color=INK_SOFT, spacing=1.2)
        y += rh + 8
    rect(s, 48, y + 2, 1184, 54, fill=AZURE_DEEP, line=None, rounded=True)
    text(s, 70, y + 13, 1140, 34,
         [("The common thread. ", {"color": GOLD_SOFT, "bold": True}),
          ("Every case runs on data that cannot leave Azerbaijan. This is precisely what a sovereign, on premise platform makes possible, and an enterprise account cannot.", {"color": WHITE})],
         size=12)
    footer(s, "Real deployments and their measured impact, adapted to national context",
           "Section 05")

    # ---------- 9 SOVEREIGN ----------
    s = new_slide(prs)
    header(s, "06", "Building the sovereign capability",
           "A sovereign capability is a full stack under national control")
    text(s, 48, 118, 500, 16, "The sovereign stack, energy to policy",
         size=11.25, color=AZURE_DEEP, bold=True, caps=True)
    layers = [
        ("7", "AI applications", "Use cases for ministries and the Administration", AZURE_DEEP),
        ("6", "Cloud & platforms", "Sovereign compute, storage and integration services", AZURE),
        ("5", "Models & data", "Models adapted on sovereign platforms, localized data", AZURE),
        ("4", "Data centers", "High efficiency data centers and GPU clusters", AZURE),
        ("3", "Connectivity", "Fast sovereign networks with redundant links", AZURE),
        ("2", "Hardware", "Locally trusted AI hardware, cooled racks", AZURE),
        ("1", "Energy & grid", "Local energy resources, a national advantage", AZURE),
        ("0", "Responsible AI", "Policy and oversight across every layer", GOLD),
    ]
    y = 142
    for idx, name, desc, col in layers:
        rect(s, 48, y, 760, 50, fill=WHITE, line=LINE, rounded=True)
        rect(s, 48, y, 50, 50, fill=col, line=None)
        text(s, 48, y + 16, 50, 20, idx, size=13.5, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
        text(s, 112, y + 16, 190, 20, name, size=12.75, color=INK, bold=True)
        text(s, 310, y + 17, 480, 20, desc, size=12, color=INK_SOFT)
        y += 56
    text(s, 838, 118, 400, 16, "Four enablers", size=11.25, color=AZURE_DEEP,
         bold=True, caps=True)
    enablers = [
        ("01", "Talent", "Recruit AI and operations skills; pair vendors with civil servants"),
        ("02", "Data & technology", "Consolidate siloed systems; label and clean data; segment networks"),
        ("03", "Governance", "Validate inputs and outputs; access by role; quarterly review"),
        ("04", "Cost", "Forecast demand; pool compute; retire cases that do not pay their way"),
    ]
    y = 142
    for num, name, desc in enablers:
        rect(s, 838, y, 394, 104, fill=WHITE, line=LINE, rounded=True)
        rect(s, 838, y, 4, 104, fill=GOLD, line=None)
        text(s, 860, y + 16, 60, 24, num, size=15, color=GOLD, bold=True)
        text(s, 900, y + 18, 310, 20, name, size=12.75, color=INK, bold=True)
        text(s, 860, y + 48, 350, 44, desc, size=11.25, color=INK_SOFT, spacing=1.25)
        y += 112
    footer(s, "Each layer depends on the one beneath it, and the whole stack stays under national control",
           "Section 06")

    # ---------- 10 WHERE TO START ----------
    s = new_slide(prs, accent=GOLD)
    header(s, "07", "Where to start",
           "Safe value now, a sovereign secure environment next")
    steps = [
        ("1", "Weeks 1 to 8", "Enterprise account live",
         "Contract it, fix data access, issue guidance on the three principles",
         "Safe drafting at scale", AZURE),
        ("2", "Quarter 1", "Capture early value",
         "Draft, translate, summarize on public and internal data. Train widely",
         "Minutes saved, measured", AZURE),
        ("3", "Quarters 1 & 2", "Accredit a secure room",
         "A secure room in a state data center, the fastest route to production",
         "Sensitive data enabled", GOLD),
        ("4", "Quarters 2 to 4", "Prove sensitive cases",
         "Citizen appeals, Security Council analysis, leadership dashboards",
         "Leadership value shown", GOLD),
        ("5", "Year 2 onward", "Scale to sovereignty",
         "Dedicated classified facility and the wider sovereign stack",
         "Full national capability", GOLD),
    ]
    rect(s, 90, 142, 1060, 3, fill=LINE_STRONG, line=None)
    for i, (num, when, title, desc, gets, col) in enumerate(steps):
        x = 48 + i * 240
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, px(x + 22), px(128), px(34), px(34))
        c.fill.solid(); c.fill.fore_color.rgb = col
        c.line.fill.background(); c.shadow.inherit = False
        text(s, x + 22, 137, 34, 20, num, size=13.5, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
        text(s, x + 22, 182, 210, 16, when, size=10.5, color=GOLD, bold=True, caps=True)
        text(s, x + 22, 202, 210, 40, title, size=13.5, color=INK, bold=True, spacing=1.15)
        text(s, x + 22, 248, 200, 70, desc, size=11.25, color=INK_SOFT, spacing=1.3)
        rect(s, x + 22, 326, 190, 1, fill=LINE, line=None)
        text(s, x + 22, 336, 200, 36,
             [("You get. ", {"color": AZURE_DEEP, "bold": True}), (gets, {})],
             size=11.25, color=INK_SOFT, spacing=1.25)
    text(s, 48, 396, 600, 16, "Prerequisites to enable this roadmap",
         size=11.25, color=AZURE_DEEP, bold=True, caps=True)
    prereqs = [
        ("Leadership mandate", "A Presidential order",
         "One accountable owner named by Sərəncam"),
        ("Data classification", "State tiers applied",
         "So every task maps to the right environment"),
        ("Enterprise agreement", "Contract in place",
         "No training on data, central control, full audit"),
        ("Funding & security path", "Budget and sign off",
         "Funding committed and state security accreditation"),
    ]
    for i, (label, head, desc) in enumerate(prereqs):
        x = 48 + i * 300
        rect(s, x, 420, 284, 180, fill=WHITE, line=LINE, rounded=True)
        rect(s, x, 420, 284, 3, fill=AZURE, line=None)
        text(s, x + 20, 440, 250, 18, label, size=11.25, color=AZURE_DEEP, bold=True)
        text(s, x + 20, 472, 250, 44, head, size=13.5, color=INK, bold=True, spacing=1.15)
        text(s, x + 20, 524, 250, 60, desc, size=11.25, color=INK_SOFT, spacing=1.3)
    footer(s, "Classification decides everything, value arrives early and compounds",
           "Section 07")

    out = "AI_Guidance_for_Leadership.pptx"
    prs.save(out)
    print("wrote", out, len(prs.slides.__iter__.__self__._sldIdLst), "slides")


if __name__ == "__main__":
    build()
